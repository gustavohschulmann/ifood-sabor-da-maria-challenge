from decimal import Decimal
from pathlib import Path
import re

from openpyxl import load_workbook

from src.domain.pantry import (
    PackageInfo,
    Pantry,
    PantryItem,
    PurchaseInfo,
    Quantity,
)


PANTRY_SHEET = "Despensa"
PRICES_SHEET = "Precos"


def to_decimal(value: object) -> Decimal:
    """
    Converts values coming from Excel into Decimal.

    Supports:
        1
        1.5
        Decimal("1.5")
        "1,5"
        "R$ 24,90"
        "24.90"
    """

    if value is None:
        raise ValueError("Expected numeric value, got None.")

    if isinstance(value, Decimal):
        return value

    if isinstance(value, int):
        return Decimal(value)

    if isinstance(value, float):
        # Never use Decimal(value) directly for floats.
        return Decimal(str(value))

    if isinstance(value, str):
        text = value.strip()

        text = text.replace("R$", "").strip()

        # Handles Brazilian formatted numbers:
        # 1.234,56 -> 1234.56
        if "," in text:
            text = text.replace(".", "")
            text = text.replace(",", ".")

        return Decimal(text)

    raise TypeError(
        f"Cannot convert {value!r} "
        f"({type(value).__name__}) to Decimal."
    )


PACKAGE_UNIT_PATTERN = re.compile(
    r"^(?P<container>un|balde)\s+"
    r"(?P<quantity>\d+(?:[.,]\d+)?)"
    r"(?P<unit>kg|g|L|l|ml)$",
    re.IGNORECASE,
)


def parse_package_unit(unit: str) -> PackageInfo | None:
    """
    Parses units such as:

        un 500g
        un 400g
        un 500ml
        balde 2kg

    Regular units such as:

        kg
        L
        un

    return None.
    """

    unit = unit.strip()

    match = PACKAGE_UNIT_PATTERN.match(unit)

    if not match:
        return None

    content_unit = match.group("unit")

    # Keep liters standardized because later conversion gets easier.
    if content_unit.lower() == "l":
        content_unit = "L"

    return PackageInfo(
        container=match.group("container").lower(),
        content_quantity=to_decimal(
            match.group("quantity")
        ),
        content_unit=content_unit,
    )


def load_pantry(file_path: str | Path) -> Pantry:
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Excel file not found: {path}"
        )

    workbook = load_workbook(
        path,
        data_only=True,
        read_only=True,
    )

    try:
        if PANTRY_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Missing sheet '{PANTRY_SHEET}'."
            )

        if PRICES_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Missing sheet '{PRICES_SHEET}'."
            )

        purchases = _load_purchases(
            workbook[PRICES_SHEET]
        )

        pantry_items: list[PantryItem] = []
        warnings: list[str] = []

        sheet = workbook[PANTRY_SHEET]

        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=2,
                values_only=True,
            ),
            start=2,
        ):
            ingredient = row[0]
            stock_quantity = row[1]
            stock_unit = row[2]

            # Skip empty rows.
            if ingredient is None:
                continue

            ingredient = str(ingredient).strip()
            stock_unit = str(stock_unit).strip()

            purchase = purchases.get(ingredient)

            if purchase is None:
                warnings.append(
                    f"No purchase/price information found "
                    f"for '{ingredient}'."
                )

            pantry_items.append(
                PantryItem(
                    ingredient=ingredient,
                    stock=Quantity(
                        value=to_decimal(
                            stock_quantity
                        ),
                        unit=stock_unit,
                    ),
                    purchase=purchase,
                    package=parse_package_unit(
                        stock_unit
                    ),
                )
            )

        return Pantry(
            items=pantry_items,
            warnings=warnings,
        )

    finally:
        workbook.close()


def _load_purchases(
    sheet,
) -> dict[str, PurchaseInfo]:

    purchases: dict[str, PurchaseInfo] = {}

    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=2,
            values_only=True,
        ),
        start=2,
    ):
        ingredient = row[0]
        purchased_quantity = row[1]
        purchased_unit = row[2]
        total_paid = row[3]

        if ingredient is None:
            continue

        ingredient = str(ingredient).strip()
        purchased_unit = str(
            purchased_unit
        ).strip()

        if ingredient in purchases:
            raise ValueError(
                f"Duplicate purchase entry for "
                f"'{ingredient}' on row "
                f"{row_number}."
            )

        purchases[ingredient] = PurchaseInfo(
            quantity=Quantity(
                value=to_decimal(
                    purchased_quantity
                ),
                unit=purchased_unit,
            ),
            total_paid=to_decimal(
                total_paid
            ),
            package=parse_package_unit(
                purchased_unit
            ),
        )

    return purchases